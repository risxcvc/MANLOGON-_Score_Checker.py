import tkinter as tk
from tkinter import messagebox
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
import os

FILE_NAME = "student_scores.xlsx"
SHEET_NAME = "Student Score Tracker"

if not os.path.exists(FILE_NAME):
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["Name", "Score", "Remarks"])
    wb.save(FILE_NAME)

def validate_inputs():
    name = name_entry.get().strip()
    score = score_entry.get().strip()

    if not name or not score:
        messagebox.showerror("Input Error", "All fields are required!")
        return False
    if not score.isdigit():
        messagebox.showerror("Input Error", "Score must be a number!")
        return False
    return True

def save_to_excel():
    if not validate_inputs():
        return

    name = name_entry.get().strip()
    score = int(score_entry.get().strip())
    remarks = "Pass" if score >= 75 else "Fail"

    wb = load_workbook(FILE_NAME)

    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb.create_sheet(SHEET_NAME)
        ws.append(["Name", "Score", "Remarks"])

    if ws.max_row > 1 and ws.cell(row=ws.max_row, column=1).value == "Average":
        ws.delete_rows(ws.max_row)

    ws.append([name, score, remarks])

    scores = [cell.value for cell in ws["B"][1:] if isinstance(cell.value, (int, float))]
    if scores:
        avg = round(sum(scores) / len(scores), 2)
        ws.append(["Average", avg, ""])

    format_excel(ws)
    wb.save(FILE_NAME)

    messagebox.showinfo("Success", "Student data saved successfully!")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)


    if ws.max_row > 1 and ws.cell(row=ws.max_row, column=1).value == "Average":
        ws.delete_rows(ws.max_row)

    ws.append([name, score, remarks])

    scores = [cell.value for cell in ws["B"][1:] if isinstance(cell.value, (int, float))]
    if scores:
        avg = round(sum(scores) / len(scores), 2)
        ws.append(["Average", avg, ""])

    format_excel(ws)
    wb.save(FILE_NAME)

    messagebox.showinfo("Success", "Student data saved successfully!")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)

def format_excel(ws):
    for cell in ws[1]:
        cell.font = Font(bold=True)

    for col in ws.columns:
        max_length = max(len(str(cell.value)) if cell.value else 0 for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max_length + 3

def show_data():
    wb = load_workbook(FILE_NAME)
    ws = wb[SHEET_NAME]

    data_window = tk.Toplevel(window)
    data_window.title("Student Records")

    col_widths = [30, 15, 15]

    for i, row in enumerate(ws.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            width = col_widths[j] if j < len(col_widths) else 15
            is_header = i == 0
            is_average = row[0] == "Average"
            is_pass_fail = (j == 2 and value in ["Pass", "Fail"])

            font_style = ("Arial", 10, "bold") if is_header or is_average or is_pass_fail else ("Arial", 10)

            label = tk.Label(
                data_window,
                text=value,
                width=width,
                font=font_style,
                borderwidth=1,
                relief="solid",
                padx=6,
                pady=3,
                bg="#FFD1DC" if not is_header else "#FF69B4",
                fg="black",
                anchor="w" if j == 0 else "center"
            )
            label.grid(row=i, column=j)

window = tk.Tk()
window.title("Student Score Tracker")

window.configure(bg="#FFE4E1")

tk.Label(window, text="Name:", bg="#FFE4E1", fg="black").grid(row=0, column=0, padx=10, pady=5, sticky="w")
tk.Label(window, text="Score:", bg="#FFE4E1", fg="black").grid(row=1, column=0, padx=10, pady=5, sticky="w")

name_entry = tk.Entry(window, width=40)
score_entry = tk.Entry(window, width=40)
name_entry.grid(row=0, column=1, padx=20, pady=5)
score_entry.grid(row=1, column=1, padx=20, pady=5)

tk.Button(window, text="Submit", command=save_to_excel, width=20, bg="#FFB6C1", fg="white").grid(row=3, column=0, columnspan=2, pady=10)
tk.Button(window, text="View Records", command=show_data, width=20, bg="#FF69B4", fg="white").grid(row=4, column=0, columnspan=2)

window.mainloop()
